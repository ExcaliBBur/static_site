"""Сборка краткого отчёта в формате Word.

Источник: docx/report.md — сжатая версия отчёта без скриншотов (полная
версия со скриншотами опубликована на сайтах, report/*.md).
Поддерживается подмножество Markdown: заголовки, абзацы, списки, таблицы,
блоки кода, изображения, **жирный**, *курсив*, `код`, ссылки.

Запуск: python scripts/build_docx.py  → report/static_site_report.docx
Затем scripts/docx_finalize.ps1 обновляет оглавление в Word и делает PDF.
"""
from __future__ import annotations

import math
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "report"
OUT = REPORT / "static_site_report.docx"
TMP = ROOT / ".cache" / "docx"
SOURCE = ROOT / "docx" / "report.md"
BODY_FONT = "Times New Roman"
CODE_FONT = "Consolas"
TEXT_WIDTH_CM = 16.5


# ---------- низкоуровневые помощники ----------

def set_cell_shading(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def shade_paragraph(par, hex_color: str) -> None:
    p_pr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    p_pr.append(shd)


def add_field(par, instr: str, placeholder: str = "") -> None:
    run = par.add_run()
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", placeholder), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
            run._r.append(el)
        elif kind == "text":
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


# ---------- inline-разметка ----------

INLINE = re.compile(
    r"(\*\*(?P<b>.+?)\*\*)|(`(?P<c>[^`]+)`)|(\*(?P<i>[^*\s][^*]*?)\*)"
    r"|(\[(?P<lt>[^\]]+)\]\((?P<lu>[^)]+)\))|(<(?P<au>https?://[^>]+)>)"
)


def add_inline(par, text: str, *, size: float | None = None, bold: bool = False, italic: bool = False) -> None:
    pos = 0

    def run(t, b=False, i=False, code=False):
        if not t:
            return
        r = par.add_run(t)
        r.bold = bold or b
        r.italic = italic or i
        if code:
            r.font.name = CODE_FONT
            r._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
            r.font.size = Pt((size or 12) - 1.5)
        elif size:
            r.font.size = Pt(size)

    for m in INLINE.finditer(text):
        run(text[pos:m.start()])
        if m.group("b"):
            run(m.group("b").replace("`", ""), b=True)
        elif m.group("c"):
            run(m.group("c"), code=True)
        elif m.group("i"):
            run(m.group("i"), i=True)
        elif m.group("lt"):
            label, url = m.group("lt"), m.group("lu")
            run(label)
            if url.startswith("http") and url.rstrip("/") not in label:
                run(f" ({url})")
        elif m.group("au"):
            run(m.group("au"))
        pos = m.end()
    run(text[pos:])


# ---------- изображения ----------

def slice_image(path: Path) -> list[tuple[Path, float]]:
    """Длинный скриншот режется на читаемые фрагменты.

    Настольный (ширина >= 1000 px): фрагменты 4:3 во всю ширину текста.
    Мобильный: фрагменты по 3 в ряд, как экраны телефона.
    """
    im = Image.open(path).convert("RGB")
    w, h = im.size
    if h / w <= 1.5:
        return [(path, TEXT_WIDTH_CM)]
    TMP.mkdir(parents=True, exist_ok=True)
    out: list[tuple[Path, float]] = []
    if w >= 1000:
        step = int(w * 0.75)
        for k in range(math.ceil(h / step)):
            part = im.crop((0, k * step, w, min(h, (k + 1) * step)))
            if part.height < 80:
                continue
            fn = TMP / f"{path.stem}_{k}.jpg"
            part.save(fn, quality=88)
            out.append((fn, TEXT_WIDTH_CM))
        return out
    cols, step, gap = 3, int(w * 1.9), 30
    n = math.ceil(h / step)
    for r in range(math.ceil(n / cols)):
        row = Image.new("RGB", (cols * w + (cols - 1) * gap, step), "white")
        for c in range(cols):
            k = r * cols + c
            if k < n:
                row.paste(im.crop((0, k * step, w, min(h, (k + 1) * step))), (c * (w + gap), 0))
        fn = TMP / f"{path.stem}_row{r}.jpg"
        row.save(fn, quality=88)
        out.append((fn, TEXT_WIDTH_CM))
    return out


# ---------- блоки ----------

class Builder:
    def __init__(self) -> None:
        self.doc = Document()
        self.fig = 0
        self.tab = 0
        self._setup()

    def _setup(self) -> None:
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin, sec.right_margin = Cm(3.0), Cm(1.5)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        st = self.doc.styles["Normal"]
        st.font.name = BODY_FONT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        st.font.size = Pt(12)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.15
        for lvl, size in ((1, 14), (2, 13), (3, 12), (4, 12)):
            hs = self.doc.styles[f"Heading {lvl}"]
            hs.font.name = BODY_FONT
            hs.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
            hs.element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
            hs.element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                hs.element.rPr.rFonts.attrib.pop(qn(attr), None)
            hs.font.size = Pt(size)
            hs.font.bold = True
            hs.font.color.rgb = RGBColor(0, 0, 0)
            hs.paragraph_format.space_before = Pt(12)
            hs.paragraph_format.space_after = Pt(6)
            hs.paragraph_format.keep_with_next = True
        sec.different_first_page_header_footer = True  # без номера на титульном листе
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(footer, "PAGE", "1")

    # --- титульный лист и оглавление ---
    def title_page(self) -> None:
        d = self.doc

        def center(text, size=14, bold=False, before=0):
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(before)
            r = p.add_run(text)
            r.font.size = Pt(size)
            r.bold = bold
            return p

        center("Федеральное государственное автономное образовательное учреждение высшего образования", 12)
        center("«Национальный исследовательский университет ИТМО»", 13, True)
        center("ОТЧЁТ", 18, True, before=150)
        center("по лабораторной работе", 14)
        center("«Генераторы статических сайтов на Python: публикация результатов экспериментов»", 15, True, before=12)
        center("Исследовательское задание T5 — публикуемость и цитируемость результата", 13, before=18)
        center("Практическое задание P2 — стресс-тест научного контента", 13)
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(120)
        p.paragraph_format.left_indent = Cm(9)
        for text, size in (
            ("Выполнил: студент группы Р4209\n", 13),
            ("Павлов Александр Сергеевич\n\n", 13),
            ("Преподаватель:\n", 13),
            ("Жуков Николай Николаевич", 13),
        ):
            p.add_run(text).font.size = Pt(size)
        center("Санкт-Петербург, 2026", 13, before=110)
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        h = d.add_paragraph()
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = h.add_run("Содержание")
        r.bold = True
        r.font.size = Pt(16)
        toc = d.add_paragraph()
        add_field(toc, 'TOC \\o "1-1" \\h \\z \\u', "Оглавление обновится при открытии документа (F9).")

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # --- элементы markdown ---
    def heading(self, text: str, level: int) -> None:
        self.doc.add_heading(re.sub(r"`", "", text), level=min(level, 4))

    def paragraph(self, text: str) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.0)
        add_inline(p, text)

    def caption(self, text: str) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, text, size=11, italic=True)

    def bullet(self, text: str, ordered: bool, level: int) -> None:
        style = "List Number" if ordered else "List Bullet"
        if level:
            style += " 2"
        p = self.doc.add_paragraph(style=style)
        add_inline(p, text)

    def code(self, lines: list[str], lang: str) -> None:
        size = 8 if max((len(x) for x in lines), default=0) > 95 else 9
        for i, line in enumerate(lines):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_after = Pt(0)
            pf.space_before = Pt(4 if i == 0 else 0)
            pf.line_spacing = 1.0
            pf.left_indent = Cm(0.3)
            shade_paragraph(p, "F3F3F1")
            r = p.add_run(line if line else " ")
            r.font.name = CODE_FONT
            r._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
            r.font.size = Pt(size)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def image(self, src: str, alt: str, base: Path) -> None:
        path = (base / src).resolve()
        if not path.exists():
            self.paragraph(f"[изображение отсутствует: {src}]")
            return
        for img_path, width in slice_image(path):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            p.add_run().add_picture(str(img_path), width=Cm(width))
        self.fig += 1

    def table(self, rows: list[list[str]]) -> None:
        ncols = len(rows[0])
        t = self.doc.add_table(rows=0, cols=ncols)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        size = 10 if ncols <= 3 else 9 if ncols <= 5 else 8
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci in range(ncols):
                txt = row[ci] if ci < len(row) else ""
                par = cells[ci].paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.0
                # Короткие таблицы не разрываются между страницами
                par.paragraph_format.keep_with_next = ri < len(rows) - 1
                if re.fullmatch(r"[\d\s,./—%≈<>+-]+", txt.replace("**", "")) and ri > 0:
                    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                add_inline(par, txt, size=size, bold=(ri == 0))
                if ri == 0:
                    set_cell_shading(cells[ci], "E8E8E4")
            if ri == 0:
                repeat_header(t.rows[0])
        # Ширина столбцов пропорциональна длине текста (корень сглаживает
        # разницу), минимум 1,4 см; сумма — ширина области текста.
        weights = [
            max(math.sqrt(max(len(r[ci]) if ci < len(r) else 0 for r in rows)), 2.2)
            for ci in range(ncols)
        ]
        total = sum(weights)
        widths = [max(1.4, TEXT_WIDTH_CM * w / total) for w in weights]
        scale = TEXT_WIDTH_CM / sum(widths)
        t.autofit = False
        for ci, w in enumerate(widths):
            for cell in t.columns[ci].cells:
                cell.width = Cm(w * scale)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # --- разбор файла ---
    def markdown(self, text: str, base: Path, first_heading_break: bool) -> None:
        lines = text.splitlines()
        i = 0
        para: list[str] = []

        def flush():
            if para:
                joined = " ".join(s.strip() for s in para)
                if re.fullmatch(r"\*[^*].*[^*]\*", joined):
                    self.caption(joined[1:-1])
                else:
                    self.paragraph(joined)
                para.clear()

        while i < len(lines):
            line = lines[i]
            s = line.strip()
            if s.startswith("<!--"):
                flush()
                while "-->" not in lines[i]:
                    i += 1
                i += 1
                continue
            if s.startswith("```"):
                flush()
                lang = s[3:].strip()
                block = []
                i += 1
                while not lines[i].strip().startswith("```"):
                    block.append(lines[i])
                    i += 1
                self.code(block, lang)
                i += 1
                continue
            m = re.match(r"^(#{1,4})\s+(.*?)\s*(\{#[^}]+\})?$", s)
            if m:
                flush()
                level = len(m.group(1))
                if level == 1 and first_heading_break:
                    self.page_break()
                self.heading(m.group(2), level)
                i += 1
                continue
            m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", s)
            if m:
                flush()
                self.image(m.group(2), m.group(1), base)
                i += 1
                continue
            if s.startswith("|"):
                flush()
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    row = lines[i].strip().strip("|")
                    cells = [c.strip() for c in re.split(r"(?<!\\)\|", row)]
                    if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                        rows.append([c.replace("\\|", "|") for c in cells])
                    i += 1
                self.table(rows)
                continue
            m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
            if m:
                flush()
                item = m.group(3)
                i += 1
                while i < len(lines) and lines[i].startswith("  ") and lines[i].strip() and not re.match(r"^\s*([-*]|\d+\.)\s", lines[i]):
                    item += " " + lines[i].strip()
                    i += 1
                self.bullet(item, m.group(2)[0].isdigit(), 1 if len(m.group(1)) >= 2 else 0)
                continue
            if not s:
                flush()
                i += 1
                continue
            para.append(line)
            i += 1
        flush()

def main() -> None:
    b = Builder()
    b.title_page()
    b.page_break()
    b.markdown(SOURCE.read_text(encoding="utf-8"), SOURCE.parent, first_heading_break=False)
    b.doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
